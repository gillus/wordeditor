"""Check the file saved by tests/e2e-select.mjs. Usage: verify_cross.py original.docx edited.docx"""
import sys

from docx import Document

from verify import RPR_ORDER, parts
from docx.oxml.ns import qn


def main(original, edited):
    orig, out = parts(original), parts(edited)
    changed = sorted(n for n in orig if orig[n] != out.get(n))
    assert changed == ["word/document.xml"] and out.keys() == orig.keys(), f"unexpected part changes: {changed}"

    doc = Document(edited)
    texts = [p.text for p in doc.paragraphs]
    assert not doc.tables, "table should have been deleted"
    for t in ["This the website for details.", "Stepge 1 of the report.", "First bullet", "Second bullet",
              "Last paragraph.Alpha", "Beta"]:
        assert t in texts, f"missing {t!r}; got {texts}"
    assert "Quarterly Report" not in texts and doc.paragraphs[0].style.name == "Normal", texts[0]

    by_text = {p.text: p for p in doc.paragraphs}
    assert [(r.text, bool(r.bold)) for r in by_text["First bullet"].runs] == [("First ", False), ("bullet", True)]
    assert [(r.text, bool(r.bold)) for r in by_text["Second bullet"].runs] == [("Second", True), (" bullet", False)]
    assert by_text["Stepge 1 of the report."].style.name == "List Number"

    xml = doc.element.body.xml
    for needle, what in [("w:hyperlink", "hyperlink"), ('w:fldCharType="begin"', "field"), ('w:name="target"', "bookmark"),
                         ("<w:ins ", "tracked change"), ("a:blip", "image"), ('w:type="page"', "page break")]:
        assert needle in xml, f"{what} lost"
    for rpr in doc.element.body.iter(qn("w:rPr")):
        ranks = [RPR_ORDER.index(c.tag.split("}")[1]) for c in rpr if c.tag.split("}")[1] in RPR_ORDER]
        assert ranks == sorted(ranks), "rPr out of order"
    print("verify_cross OK")


if __name__ == "__main__":
    main(*sys.argv[1:])
