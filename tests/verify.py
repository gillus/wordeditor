"""Check the documents produced by tests/e2e.mjs: usage: verify.py original.docx pristine.docx edited.docx"""
import sys
import zipfile

from docx import Document
from docx.oxml.ns import qn

RPR_ORDER = [
    "rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike", "outline",
    "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden", "color", "spacing",
    "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText",
    "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath", "rPrChange",
]


def parts(path):
    with zipfile.ZipFile(path) as z:
        return {name: z.read(name) for name in z.namelist()}


def main(original, pristine, edited):
    orig = parts(original)
    assert parts(pristine) == orig, "unedited save changed the package"

    out = parts(edited)
    assert out.keys() == orig.keys(), "package parts added or removed"
    changed = sorted(n for n in orig if orig[n] != out[n])
    assert changed == ["word/document.xml"], f"unexpected parts changed: {changed}"

    doc = Document(edited)
    texts = [p.text for p in doc.paragraphs]
    expect = [
        "This is bold and italic text. Appended.",
        "See the website for details.!",
        "Bookmarked",
        " paragraph.",
        "First bulletSecond bullet",
        "Step one\ncontinued",
        "Step two",
        "Before the break.",
        "Last paragraph.Pasted A",
        "Pasted B",
    ]
    for t in expect:
        assert t in texts, f"missing paragraph {t!r}; got {texts}"

    body = doc.element.body
    xml = body.xml
    assert 'w:hyperlink' in xml and "the website" in xml, "hyperlink lost"
    assert 'w:name="target"' in xml and "<w:bookmarkEnd" in xml, "bookmark lost"
    assert 'w:fldCharType="begin"' in xml and " PAGE " in xml, "field lost"
    assert "<w:ins " in xml and "inserted by a reviewer" in xml, "tracked change lost"
    assert 'w:type="page"' in xml, "page break lost"
    assert "a:blip" in xml, "image lost"

    by_text = {p.text: p for p in doc.paragraphs}
    last = by_text["Last paragraph.Pasted A"]
    assert last.runs[0].text == "Last" and last.runs[0].bold, "bold toggle not saved"
    assert not last.runs[1].bold

    intro = by_text["This is bold and italic text. Appended."]
    assert not any(r.bold for r in intro.runs), "un-bolding not saved"
    assert [r.text for r in intro.runs if r.italic] == ["italic"]

    pasted = by_text["Pasted B"]
    assert pasted.runs[0].text == "Pasted" and pasted.runs[0].italic and pasted.runs[0].underline
    assert not pasted.runs[1].italic and not pasted.runs[1].underline

    assert by_text["Before the break."].style.style_id == "Heading2"
    assert by_text["Pasted B"].style.style_id == by_text["Last paragraph.Pasted A"].style.style_id

    cell = doc.tables[0].cell(2, 2)
    assert cell.text == "R2C2\tX", repr(cell.text)

    # Numbering survives on the merged list item.
    merged = by_text["First bulletSecond bullet"]
    assert merged.style.name == "List Bullet"

    # Every rPr must keep schema order, or Word reports the file as corrupt.
    for rpr in body.iter(qn("w:rPr")):
        names = [c.tag.split("}")[1] for c in rpr]
        ranks = [RPR_ORDER.index(n) for n in names if n in RPR_ORDER]
        assert ranks == sorted(ranks), f"rPr out of order: {names}"

    print("verify OK")


if __name__ == "__main__":
    main(*sys.argv[1:])
