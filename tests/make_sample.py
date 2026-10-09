"""Generate tests/fixtures/sample.docx covering the features the editor must render and preserve."""
import struct
import zlib
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches

OUT = Path(__file__).parent / "fixtures" / "sample.docx"


def tiny_png(path, w=40, h=20):
    raw = b"".join(b"\x00" + bytes([200, 80, 60]) * w for _ in range(h))

    def chunk(tag, data):
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data))

    path.write_bytes(
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


def add_hyperlink(paragraph, url, text):
    r_id = paragraph.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True
    )
    paragraph._p.append(parse_xml(
        f'<w:hyperlink {nsdecls("w", "r")} r:id="{r_id}"><w:r><w:rPr><w:rStyle w:val="Hyperlink"/></w:rPr>'
        f"<w:t>{text}</w:t></w:r></w:hyperlink>"
    ))


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading("Quarterly Report", level=1)

    p = doc.add_paragraph("This is ")
    p.add_run("bold").bold = True
    p.add_run(" and ")
    p.add_run("italic").italic = True
    p.add_run(" text.")

    p = doc.add_paragraph("See ")
    add_hyperlink(p, "https://example.com", "the website")
    p.add_run(" for details.")

    p = doc.add_paragraph("Bookmarked paragraph.")
    p._p.insert(1, parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="7" w:name="target"/>'))
    p._p.append(parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="7"/>'))

    doc.add_paragraph("First bullet", style="List Bullet")
    doc.add_paragraph("Second bullet", style="List Bullet")
    doc.add_paragraph("Step one", style="List Number")
    doc.add_paragraph("Step two", style="List Number")

    table = doc.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            cell.text = f"R{r}C{c}"
    table.cell(0, 0).merge(table.cell(0, 1))

    p = doc.add_paragraph("Page ")
    p._p.append(parse_xml(
        f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="begin"/></w:r>'))
    p._p.append(parse_xml(
        f'<w:r {nsdecls("w")}><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r>'))
    p._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="separate"/></w:r>'))
    p._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:t>1</w:t></w:r>'))
    p._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:fldChar w:fldCharType="end"/></w:r>'))
    p.add_run(" of the report.")

    p = doc.add_paragraph("Reviewed text ")
    p._p.append(parse_xml(
        f'<w:ins {nsdecls("w")} w:id="1" w:author="Reviewer" w:date="2026-01-01T00:00:00Z">'
        f"<w:r><w:t>inserted by a reviewer</w:t></w:r></w:ins>"
    ))

    img = OUT.parent / "pixel.png"
    tiny_png(img)
    doc.add_paragraph().add_run().add_picture(str(img), width=Inches(1))
    img.unlink()

    doc.add_paragraph("Before the break.").add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph("Last paragraph.")
    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
